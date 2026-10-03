import os
import time
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
import win32com.client as win32

# =================CONFIGURATION=================
EXCEL_FILE = "recipients.xlsx"
GLOBAL_DEFAULT_ACTION = "DRAFT"  # "DRAFT" or "SEND"
OVERRIDE_FORCE_DRAFT = True      # Keep True for testing; False for direct live dispatch
DELAY_BETWEEN_EMAILS = 1.2       # Seconds pause between iterations
SKIP_ALREADY_SENT = True         # Skips rows already marked 'Sent'
# ===============================================

# Status styling
FILL_SENT = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FILL_DRAFTED = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FILL_FAILED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

FONT_SENT = Font(color="006100", bold=True)
FONT_DRAFTED = Font(color="9C6500", bold=True)
FONT_FAILED = Font(color="9C0006", bold=True)


def apply_status_style(cell, status_text: str):
    s = status_text.lower()
    if s == "sent":
        cell.fill, cell.font = FILL_SENT, FONT_SENT
    elif s == "drafted":
        cell.fill, cell.font = FILL_DRAFTED, FONT_DRAFTED
    elif "failed" in s:
        cell.fill, cell.font = FILL_FAILED, FONT_FAILED


def resolve_attachments(raw_entry: str) -> list[str]:
    """
    Parses single files, semicolon-delimited lists, or directories.
    Returns a list of validated absolute paths.
    """
    if not raw_entry or not str(raw_entry).strip():
        return []

    collected = []
    # Split by semicolon or comma to support lists
    entries = str(raw_entry).replace(",", ";").split(";")

    for item in entries:
        cleaned = item.strip().strip('"').strip("'")
        if not cleaned:
            continue

        abs_p = os.path.abspath(cleaned)

        if os.path.isdir(abs_p):
            # If entry is a folder, attach all non-temporary files inside it
            for r, _, fs in os.walk(abs_p):
                for f in fs:
                    if not f.startswith("~$") and not f.startswith("."):
                        collected.append(os.path.join(r, f))
        elif os.path.isfile(abs_p):
            collected.append(abs_p)
        else:
            print(f"  [Warning] Attachment not found: {abs_p}")

    return list(dict.fromkeys(collected))


def format_custom_html_body(raw_body: str, name: str, company: str) -> str:
    """
    Builds the HTML message body. If raw_body is provided from Excel,
    it wraps it cleanly and converts line breaks into HTML tags.
    If empty, it falls back to a default template.
    """
    if raw_body and raw_body.strip():
        # Convert Excel newlines (Alt+Enter) to HTML breaks
        formatted_paragraphs = raw_body.strip().replace("\r\n", "\n").replace("\n", "<br>")
        content = f"<p>{formatted_paragraphs}</p>"
    else:
        # Default fallback message
        content = f"""
        <p>Hello <strong>{name}</strong>,</p>
        <p>Please find attached the latest documents and updates regarding <strong>{company}</strong>.</p>
        <p>Feel free to reach out if you have any questions or require adjustments.</p>
        """

    return f"""
    <html>
        <body style="font-family: Calibri, Arial, sans-serif; font-size: 11pt; color: #333333; line-height: 1.5;">
            {content}
            <br>
        </body>
    </html>
    """


def run_bulk_emailer():
    if not os.path.exists(EXCEL_FILE):
        print(f"Error: '{EXCEL_FILE}' not found.")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)
    sheet = wb.active

    # Dynamically map all column header names to column index numbers
    header_col_map = {}
    for c in range(1, sheet.max_column + 1):
        v = sheet.cell(row=1, column=c).value
        if v:
            header_col_map[str(v).strip()] = c

    # Ensure tracking columns exist
    max_c = sheet.max_column
    if "Status" not in header_col_map:
        max_c += 1
        sheet.cell(row=1, column=max_c, value="Status")
        header_col_map["Status"] = max_c
    if "Timestamp" not in header_col_map:
        max_c += 1
        sheet.cell(row=1, column=max_c, value="Timestamp")
        header_col_map["Timestamp"] = max_c

    try:
        outlook = win32.Dispatch("Outlook.Application")
    except Exception as e:
        print(f"Error connecting to Outlook: {e}")
        print("Please verify that Microsoft Outlook is open and running.")
        return

    print("--- Bulk Email Automation Started ---")
    print(f"Force Draft Safety: {OVERRIDE_FORCE_DRAFT}\n")

    for row_idx in range(2, sheet.max_row + 1):
        def get_val(h):
            c = header_col_map.get(h)
            v = sheet.cell(row=row_idx, column=c).value if c else None
            return "" if v is None else str(v).strip()

        to_email = get_val("Email")
        if not to_email:
            continue

        curr_status = get_val("Status")
        if SKIP_ALREADY_SENT and curr_status.lower() == "sent":
            print(f"Row {row_idx}: Skipped (Already Sent to {to_email}).")
            continue

        cc_email = get_val("CC")
        bcc_email = get_val("BCC")
        row_subject = get_val("Subject")
        row_body = get_val("Body")
        name = get_val("Name") or "Valued Client"
        company = get_val("Company") or "your account"
        raw_attach = get_val("Attachments")

        # Subject logic
        final_subject = row_subject if row_subject else f"Important Update for {company}"

        # Action logic
        row_action = get_val("Action").upper()
        target_action = row_action if row_action in ["SEND", "DRAFT"] else GLOBAL_DEFAULT_ACTION.upper()
        if OVERRIDE_FORCE_DRAFT:
            target_action = "DRAFT"

        status_result = ""

        try:
            # 1. Initialize item and retrieve default signature
            mail = outlook.CreateItem(0)  # 0 = olMailItem
            mail.Display()                # Loads user signature into HTMLBody
            signature = mail.HTMLBody

            # 2. Assign recipients and metadata
            mail.To = to_email
            if cc_email:
                mail.CC = cc_email
            if bcc_email:
                mail.BCC = bcc_email

            mail.Subject = final_subject

            # 3. Assemble custom body + signature
            custom_html = format_custom_html_body(row_body, name, company)
            mail.HTMLBody = custom_html + signature

            # 4. Attach all resolved files
            files = resolve_attachments(raw_attach)
            for f in files:
                mail.Attachments.Add(f)

            # 5. Dispatch or Draft
            if target_action == "SEND":
                mail.Send()
                status_result = "Sent"
                print(f"Row {row_idx}: [SENT] -> {to_email} | Subject: '{final_subject}' ({len(files)} files)")
            else:
                mail.Save()
                mail.Close(0)  # 0 = olSave (saves draft and closes window)
                status_result = "Drafted"
                print(f"Row {row_idx}: [DRAFTED] -> {to_email} | Subject: '{final_subject}' ({len(files)} files)")

        except Exception as err:
            status_result = "Failed"
            print(f"Row {row_idx}: [ERROR] {to_email} -> {err}")

        # 6. Update Excel audit status
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        status_cell = sheet.cell(row=row_idx, column=header_col_map["Status"])
        status_cell.value = status_result
        apply_status_style(status_cell, status_result)

        sheet.cell(row=row_idx, column=header_col_map["Timestamp"], value=now_str)
        wb.save(EXCEL_FILE)

        time.sleep(DELAY_BETWEEN_EMAILS)

    print("\nProcessing complete. All rows updated in Excel.")


if __name__ == "__main__":
    run_bulk_emailer()