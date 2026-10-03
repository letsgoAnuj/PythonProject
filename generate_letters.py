import os
import re
import zipfile
import docx
import openpyxl
from tqdm import tqdm

# 1. Automatic File Path Resolution
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


# Match files flexibly regardless of slight naming differences
def find_file(patterns):
    for f in os.listdir(SCRIPT_DIR):
        for p in patterns:
            if p.lower() in f.lower() and not f.startswith("~$") and not f.startswith("Letter_"):
                return os.path.join(SCRIPT_DIR, f)
    return None


EXCEL_FILE = find_file(["Vendor Performace rating sheet", "Vendor Performance rating sheet"])
TEMPLATE_DOCX = find_file(["POLYCOATERS", "Vendor_Performance_Rating.docx"])

if not EXCEL_FILE:
    raise FileNotFoundError(
        "Could not find Excel file ('Vendor Performace rating sheet.xlsx') in the project directory.")
if not TEMPLATE_DOCX:
    raise FileNotFoundError("Could not find template Word file (.docx) in the project directory.")

DOCX_DIR = os.path.join(SCRIPT_DIR, "Generated_Letters_DOCX")
PDF_DIR = os.path.join(SCRIPT_DIR, "Generated_Letters_PDF")
os.makedirs(DOCX_DIR, exist_ok=True)
os.makedirs(PDF_DIR, exist_ok=True)

print(f"Using Excel: {os.path.basename(EXCEL_FILE)}")
print(f"Using Template: {os.path.basename(TEMPLATE_DOCX)}")


def clean_filename(name):
    """Sanitize vendor name for valid Windows/Linux file paths."""
    return re.sub(r'[\\/*?:"<>|]', "", str(name)).strip()


def extract_vendor_data(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Locate scorecard sheet case-insensitively
    ws = None
    for sheet_name in wb.sheetnames:
        if "scorecard" in sheet_name.lower():
            ws = wb[sheet_name]
            break
    if ws is None:
        ws = wb.active

    vendor_data = []
    max_row = ws.max_row
    max_col = ws.max_column

    # Scan grid for 'Business Partner' cells
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            val = str(ws.cell(row=r, column=c).value or "").strip()
            if "business partner" in val.lower():
                vname = ws.cell(row=r, column=c + 1).value
                if not vname:
                    continue
                vname = str(vname).strip()

                scores = {}
                # Extract the scorecard rows downstream
                for sub_r in range(r + 1, min(r + 10, max_row + 1)):
                    p_name = str(ws.cell(row=sub_r, column=c).value or "").strip()
                    s_val = ws.cell(row=sub_r, column=c + 2).value
                    if s_val is None:
                        s_val = ws.cell(row=sub_r, column=c + 1).value

                    if s_val is not None and p_name:
                        try:
                            num = float(s_val)
                            # Convert 0.30 -> 30, 0.73 -> 73.00 if stored as decimals
                            if num <= 1.0 and "weightage" not in p_name.lower():
                                num = num * 100
                            if "overall" in p_name.lower():
                                formatted = f"{num:.2f}"
                            else:
                                formatted = f"{int(round(num))}" if round(num, 2).is_integer() else f"{num:.2f}"
                        except (ValueError, TypeError):
                            formatted = str(s_val).strip()

                        if "delivery" in p_name.lower():
                            scores["Delivery"] = formatted
                        elif "quality" in p_name.lower():
                            scores["Quality (Rejection)"] = formatted
                        elif "testing" in p_name.lower():
                            scores["Testing Compliance"] = formatted
                        elif "response" in p_name.lower():
                            scores["Response"] = formatted
                        elif "overall" in p_name.lower():
                            scores["Overall Rating"] = formatted

                if scores:
                    vendor_data.append({"vendor_name": vname, "scores": scores})
                break

    return vendor_data


def generate_letter(template_path, out_path, vendor_name, scores):
    """Fills template in-place: preserves headers, footers, watermarks, margins & 1-page fit."""
    doc = docx.Document(template_path)

    # 1. Update Vendor Name in address block
    for p in doc.paragraphs:
        if "A C POLYCOATERS PRIVATE LIMITED" in p.text:
            for run in p.runs:
                if "A C POLYCOATERS PRIVATE LIMITED" in run.text:
                    run.text = run.text.replace("A C POLYCOATERS PRIVATE LIMITED", vendor_name)
        elif "[VENDOR NAME]" in p.text:
            for run in p.runs:
                if "[VENDOR NAME]" in run.text:
                    run.text = run.text.replace("[VENDOR NAME]", vendor_name)

    # 2. Update Table Scores
    for table in doc.tables:
        for row in table.rows:
            header_text = row.cells[0].text.strip()
            for param, score_val in scores.items():
                if param.lower() in header_text.lower():
                    target_cell = row.cells[2]
                    # Preserve existing run font styling
                    if target_cell.paragraphs and target_cell.paragraphs[0].runs:
                        target_cell.paragraphs[0].runs[0].text = str(score_val)
                    else:
                        target_cell.text = str(score_val)

    doc.save(out_path)


def main():
    print("\nExtracting vendor records from Excel...")
    vendors = extract_vendor_data(EXCEL_FILE)
    print(f"Found {len(vendors)} vendor scorecards.\n")

    if not vendors:
        print("No vendors found. Please verify the sheet contents.")
        return

    generated_docx_files = []
    for item in tqdm(vendors, desc="Generating DOCX Letters"):
        vname = item["vendor_name"]
        safe_name = clean_filename(vname)
        filename = f"Letter_{safe_name}_Vendor_Performance_Rating.docx"
        out_path = os.path.join(DOCX_DIR, filename)

        generate_letter(TEMPLATE_DOCX, out_path, vname, item["scores"])
        generated_docx_files.append(out_path)

    # Convert to PDF
    print("\nConverting documents to PDF (preserving single-page letterhead layout)...")
    try:
        from docx2pdf import convert
        convert(DOCX_DIR, PDF_DIR)
        print("PDF conversion completed successfully.")
    except Exception as e:
        print(f"Note: PDF conversion requires MS Word installed: {e}")

    # Create ZIP files
    print("\nCreating ZIP packages...")
    zip_docx_path = os.path.join(SCRIPT_DIR, "Vendor_Letters_DOCX.zip")
    zip_pdf_path = os.path.join(SCRIPT_DIR, "Vendor_Letters_PDF.zip")

    with zipfile.ZipFile(zip_docx_path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in os.listdir(DOCX_DIR):
            if f.endswith(".docx"):
                z.write(os.path.join(DOCX_DIR, f), arcname=f)

    with zipfile.ZipFile(zip_pdf_path, "w", zipfile.ZIP_DEFLATED) as z:
        for f in os.listdir(PDF_DIR):
            if f.endswith(".pdf"):
                z.write(os.path.join(PDF_DIR, f), arcname=f)

    print("\nProcess finished successfully!")
    print(f"1. {zip_docx_path}")
    print(f"2. {zip_pdf_path}")


if __name__ == "__main__":
    main()