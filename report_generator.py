import pandas as pd
import matplotlib.pyplot as plt
import os
import time
import warnings
import win32com.client as win32
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import matplotlib.dates as mdates
from matplotlib.backends.backend_pdf import PdfPages
import matplotlib.patches as patches
from datetime import datetime
import sys

sys.stdout.reconfigure(encoding='utf-8')
warnings.filterwarnings('ignore')

# Set up paths
BASE_DIR = os.getcwd()
ANALYSIS_DIR = os.path.join(BASE_DIR, "analysis_output")
CHART_DIR = os.path.join(ANALYSIS_DIR, "charts")
DATA_FILE = os.path.join(ANALYSIS_DIR, "Master_Trend_Data.xlsx")
PPT_OUTPUT = os.path.join(ANALYSIS_DIR, "Footwear_Procurement_Forecast.pptx")
PDF_OUTPUT = os.path.join(ANALYSIS_DIR, "Footwear_Procurement_Forecast.pdf")

if not os.path.exists(CHART_DIR):
    os.makedirs(CHART_DIR)


# ==========================================
# 0. PRE-FLIGHT & POST-FLIGHT ROBUSTNESS
# ==========================================
def auto_save_and_close_ppt(target_filepath):
    """Detects if the target PPTX file is open, saves it, and closes it gracefully."""
    abs_path = os.path.abspath(target_filepath)
    try:
        ppt = win32.GetObject(None, "PowerPoint.Application")
        for pres in ppt.Presentations:
            if pres.FullName.lower() == abs_path.lower():
                print(f"⚠️ Detected {os.path.basename(abs_path)} is open. Auto-saving and closing...")
                pres.Save()
                pres.Close()
                print("✓ Presentation closed successfully.")
                break
    except Exception:
        pass


def pre_flight_checks():
    print("\n--- Phase 0: Pre-Flight Lock Checks ---")
    auto_save_and_close_ppt(PPT_OUTPUT)

    # PDF Lock check (Adobe Acrobat Trap)
    if os.path.exists(PDF_OUTPUT):
        try:
            with open(PDF_OUTPUT, "a"):
                pass
        except PermissionError:
            error_msg = f"CRITICAL ERROR: {os.path.basename(PDF_OUTPUT)} is locked (likely open in Adobe Acrobat). Please close it."
            print(error_msg, file=sys.stderr)
            sys.exit(1)  # Triggers automator admin trap
    print("✓ Output targets are unlocked and ready.")


def post_flight_validation():
    """Post-Flight: The Stale File Validation"""
    print("\n--- Phase 3: Post-Flight Validation ---")
    current_time = time.time()
    stale_threshold = 300  # 5 minutes

    for filepath in [PPT_OUTPUT, PDF_OUTPUT]:
        if os.path.exists(filepath):
            mtime = os.path.getmtime(filepath)
            if current_time - mtime > stale_threshold:
                error_msg = f"CRITICAL ERROR: {os.path.basename(filepath)} was not updated (Stale file detected)."
                print(error_msg, file=sys.stderr)
                sys.exit(1)
        else:
            error_msg = f"CRITICAL ERROR: {os.path.basename(filepath)} was not generated."
            print(error_msg, file=sys.stderr)
            sys.exit(1)
    print("✓ Outputs validated successfully. No stale files detected.")


# ==========================================
# 1. EXPANDED SLIDE CONFIGURATION
# ==========================================
PAGES_CONFIG = [
    {
        "title": "Macro Impact: USD Exchange Tracker",
        "insight": "FORECAST: Tracking USD fluctuations to gauge general import cost inflation and macro-level procurement impacts on imported components.",
        "charts": [
            {"title": "USD Trend vs INR", "products": ["USD-INR"], "sheet": "Exchange_Rates", "price_col": "Price",
             "unit": "INR"}
        ]
    },
    {
        "title": "Macro Impact: EUR Exchange Tracker",
        "insight": "FORECAST: Tracking EUR fluctuations to monitor European import costs, specifically for specialized hardware and premium synthetics.",
        "charts": [
            {"title": "EUR Trend vs INR", "products": ["EUR-INR"], "sheet": "Exchange_Rates", "price_col": "Price",
             "unit": "INR"}
        ]
    },
    {
        "title": "Regional Impact: RMB Exchange Tracker",
        "insight": "FORECAST: Tracking the Chinese Yuan (RMB) to INR trajectory. Direct impacts on landed costs for imported Chinese components, including synthetics, hardware, and specialized textiles.",
        "charts": [
            {"title": "CNY (RMB) Trend vs INR", "products": ["CNY-INR"], "sheet": "Exchange_Rates",
             "price_col": "Price", "unit": "INR"}
        ]
    },
    {
        "title": "Upstream Drivers: Items Impacting Yarn",
        "insight": "FORECAST: Track PTA, MEG, and PP for cascading price adjustments in knitted vamps, airmesh, and woven labels.",
        "charts": [
            {"title": "PTA & MEG Trajectory", "products": ["PTA", "MEG"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"},
            {"title": "PP Trend", "products": ["PP"], "sheet": "Product_Trends", "price_col": "Final_INR_KG_Price",
             "unit": "INR/KG"}
        ]
    },
    {
        "title": "Global Upstream: Brent Crude",
        "insight": "FORECAST: Brent Crude acts as the ultimate upstream driver, heavily dictating the baseline costs for all downstream petrochemical derivatives.",
        "charts": [
            {"title": "Brent Crude Trajectory", "products": ["Brent"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "USD/BBL"}
        ]
    },
    {
        "title": "Coated Fabrics: Items Impacting PVC",
        "insight": "FORECAST: DOP (plasticizer) remains highly volatile. While PVC Suspension offers a stable baseline for fabrics, DOP drives the cost of flexibility.",
        "charts": [
            {"title": "DOP & PVC Suspension Trend", "products": ["DOP", "PVC_Suspension"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"}
        ]
    },
    {
        "title": "Packaging & Films: LDPE & LLDPE",
        "insight": "FORECAST: Tracking Low-Density and Linear Low-Density Polyethylene. These commodities drive the baseline costs for flexible packaging films and protective wraps.",
        "charts": [
            {"title": "LDPE & LLDPE Trajectory", "products": ["LDPE", "LLDPE"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"}
        ]
    },
    {
        "title": "Raw Precursors: Naphtha & Acetone",
        "insight": "FORECAST: Naphtha and Acetone act as primary precursors. Spikes in these upstream commodities will cascade into PU foam and EVA sheet cost increases.",
        "charts": [
            {"title": "Naphtha Trend", "products": ["Naphtha"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"},
            {"title": "Acetone Trend", "products": ["ACETONE"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"}
        ]
    },
    {
        "title": "Sole Materials: EVA & Latex",
        "insight": "ANALYSIS: Divergence between natural Latex and synthetic EVA allows for strategic recipe blending.",
        "charts": [
            {"title": "EVA Trend", "products": ["EVA"], "sheet": "Product_Trends", "price_col": "Final_INR_KG_Price",
             "unit": "INR/KG"},
            {"title": "Latex Trend", "products": ["LATEX"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"}
        ]
    },
    {
        "title": "Hardware Trims: Brass & Zinc",
        "insight": "ANALYSIS: Zinc (galvanizing) and Brass dictate hardware/accessory costs. Watch global LME trends to prevent domestic suppliers from inflating pricing.",
        "charts": [
            {"title": "Brass Trend", "products": ["BRASS"], "sheet": "Product_Trends",
             "price_col": "Final_INR_KG_Price", "unit": "INR/KG"},
            {"title": "Zinc Trend", "products": ["ZINC"], "sheet": "Product_Trends", "price_col": "Final_INR_KG_Price",
             "unit": "INR/KG"}
        ]
    }
]

RESOURCE_LINKS = [
    ("PolymerUpdate", "PVC, EVA, Naphtha, DOP, Acetone, PET, PP, Brent Crude, LDPE, LLDPE"),
    ("RBI (Reserve Bank of India)", "USD-INR, EUR-INR Exchange Rates"),
    ("ExchangeRates.org.uk", "CNY-INR (Chinese Yuan) Exchange Rates"),
    ("Westmetall", "Zinc (LME), Brass (Domestic/Imported)"),
    ("PCK Ltd", "Natural Rubber / Latex")
]

COLORS = ['#0052cc', '#e52d27', '#00b8d9', '#6554c0', '#ff8b00']


# ==========================================
# 2. POWERPOINT GENERATOR ENGINE
# ==========================================
def create_slide_image(page_config, df_products, df_fx):
    num_charts = len(page_config["charts"])
    fig, axes = plt.subplots(1, num_charts, figsize=(12, 4.2))
    fig.patch.set_facecolor('white')

    if num_charts == 1:
        axes = [axes]

    one_year_ago = pd.to_datetime('today') - pd.DateOffset(years=1)

    for ax, chart_config in zip(axes, page_config["charts"]):
        ax.set_facecolor('white')
        y_max, y_min = 0, float('inf')
        target_df = df_fx if chart_config["sheet"] == "Exchange_Rates" else df_products
        unit = chart_config["unit"]

        if not target_df.empty and 'Date' in target_df.columns:
            df_1yr = target_df[target_df['Date'] >= one_year_ago]
            for idx, prod in enumerate(chart_config["products"]):
                mat_data = df_1yr[df_1yr['Product'] == prod].sort_values('Date')
                if not mat_data.empty:
                    price_col = chart_config["price_col"]
                    monthly_data = mat_data.resample('MS', on='Date').mean(numeric_only=True).reset_index().dropna(
                        subset=[price_col])

                    if monthly_data.empty: continue
                    if monthly_data[price_col].max() > y_max: y_max = monthly_data[price_col].max()
                    if monthly_data[price_col].min() < y_min: y_min = monthly_data[price_col].min()

                    color = COLORS[idx % len(COLORS)]
                    ax.plot(monthly_data['Date'], monthly_data[price_col], label=prod, linewidth=3.0, marker='o',
                            markersize=7, markeredgecolor='white', markeredgewidth=1.5, color=color, zorder=3)

                    # Dynamic Prefix based on Unit
                    prefix = "$" if unit == "USD/BBL" else ("" if unit == "INR" else "₹")

                    for x, y in zip(monthly_data['Date'], monthly_data[price_col]):
                        offset = (y_max - y_min) * 0.06 if y_max != y_min else y_max * 0.06

                        format_str = f"{prefix}{y:,.2f}" if unit in ["USD/BBL", "INR"] else f"{prefix}{y:,.0f}"

                        ax.text(x, y + offset, format_str, fontsize=9.5, ha='center', va='bottom', color=color,
                                fontweight='bold', zorder=4,
                                bbox=dict(facecolor='white', edgecolor='none', alpha=0.75, pad=1.5))

        ax.set_title(chart_config["title"], fontsize=12, fontweight='bold', color='#1a1a1a', pad=15)
        ax.set_ylabel(f"{unit}", fontsize=10, color='#4a4a4a', labelpad=10)

        # Graceful Degradation: Draw watermark if data is empty
        if y_min == float('inf'):
            ax.set_ylim(bottom=0, top=100)
            ax.text(0.5, 0.5, "Data Unavailable", ha='center', va='center', transform=ax.transAxes, color='#e52d27',
                    fontsize=12, fontweight='bold')
        else:
            ax.set_ylim(bottom=y_min * 0.85, top=y_max * 1.28)

        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['left'].set_color('#e0e0e0')
        ax.spines['bottom'].set_color('#e0e0e0')
        ax.grid(axis='y', linestyle='-', color='#f0f0f0')
        ax.xaxis.set_major_locator(mdates.MonthLocator(interval=1))
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%b %y'))
        ax.tick_params(axis='x', rotation=45, labelsize=9.5)
        ax.legend(loc="upper left", fontsize=9, frameon=False)

    plt.tight_layout()
    filename = page_config['title'].replace(" ", "_").replace(",", "").replace("&", "and")[:25] + ".png"
    chart_path = os.path.join(CHART_DIR, filename)
    plt.savefig(chart_path, dpi=400, bbox_inches='tight', facecolor='white')
    plt.close()
    return chart_path


def add_corporate_header(slide, title_text):
    tx_title = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.8))
    p = tx_title.text_frame.add_paragraph()
    p.text = title_text
    p.font.size = Pt(26)
    p.font.color.rgb = RGBColor(15, 23, 42)
    p.font.bold = True

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(1.0), Inches(12.533), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 82, 204)
    line.line.color.rgb = RGBColor(0, 82, 204)


def build_presentation():
    print("\nGenerating PPTX presentation...")

    # Safely load dataframes so a single missing sheet doesn't crash everything
    try:
        df_products = pd.read_excel(DATA_FILE, sheet_name="Product_Trends")
        df_products['Date'] = pd.to_datetime(df_products['Date'])
    except Exception as e:
        print(f"Warning: Could not load Product_Trends: {e}")
        df_products = pd.DataFrame()

    try:
        df_fx = pd.read_excel(DATA_FILE, sheet_name="Exchange_Rates")
        df_fx['Date'] = pd.to_datetime(df_fx['Date'])
    except Exception as e:
        print(f"Warning: Could not load Exchange_Rates: {e}")
        df_fx = pd.DataFrame()

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1. Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg.line.color.rgb = RGBColor(15, 23, 42)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.4))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(0, 82, 204)
    accent.line.color.rgb = RGBColor(0, 82, 204)

    tx_brand = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.333), Inches(0.8))
    p1 = tx_brand.text_frame.add_paragraph()
    p1.text = "C A M P U S   A C T I V E W E A R"
    p1.font.size, p1.font.color.rgb, p1.font.bold = Pt(18), RGBColor(148, 163, 184), True
    p1.alignment = PP_ALIGN.CENTER

    tx_title = slide.shapes.add_textbox(Inches(1), Inches(2.9), Inches(11.333), Inches(1.5))
    p2 = tx_title.text_frame.add_paragraph()
    p2.text = "Procurement & Sourcing"
    p2.font.size, p2.font.color.rgb, p2.font.bold = Pt(54), RGBColor(255, 255, 255), True
    p2.alignment = PP_ALIGN.CENTER

    tx_sub = slide.shapes.add_textbox(Inches(1), Inches(4.1), Inches(11.333), Inches(1.0))
    p3 = tx_sub.text_frame.add_paragraph()
    p3.text = "Raw Material Cost & Market Analytics"
    p3.font.size, p3.font.color.rgb = Pt(28), RGBColor(226, 232, 240)
    p3.alignment = PP_ALIGN.CENTER

    tx_footer = slide.shapes.add_textbox(Inches(1), Inches(6.0), Inches(11.333), Inches(0.8))
    p4 = tx_footer.text_frame.add_paragraph()
    p4.text = f"Automated 1-Year Market Intelligence Pipeline  |  Generated: {datetime.now().strftime('%d %b %Y')}"
    p4.font.size, p4.font.color.rgb, p4.font.italic = Pt(14), RGBColor(148, 163, 184), True
    p4.alignment = PP_ALIGN.CENTER

    # 2. Matrix Slides
    try:
        matrix_df = pd.read_excel(DATA_FILE, sheet_name="Variation_Matrix").iloc[:, :11]
        mid_point = (len(matrix_df) + 1) // 2
        parts = [(matrix_df.iloc[:mid_point], "Part 1"), (matrix_df.iloc[mid_point:], "Part 2")]

        for part_df, part_label in parts:
            if part_df.empty: continue

            slide = prs.slides.add_slide(prs.slide_layouts[6])
            add_corporate_header(slide, f"Executive Summary: Price Variation Matrix ({part_label})")

            rows, cols = len(part_df) + 1, len(part_df.columns)
            left, top = Inches(0.4), Inches(1.6)
            width, height = Inches(12.533), Inches(0.4 * rows)

            table_shape = slide.shapes.add_table(rows, cols, left, top, width, height).table

            for col_idx, col_name in enumerate(part_df.columns):
                cell = table_shape.cell(0, col_idx)
                cell.text = str(col_name)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0, 82, 204)
                for p in cell.text_frame.paragraphs:
                    p.font.bold, p.font.color.rgb, p.font.size = True, RGBColor(255, 255, 255), Pt(11)
                    p.alignment = PP_ALIGN.CENTER

            for r_idx, (_, row) in enumerate(part_df.iterrows()):
                row_bg = RGBColor(248, 250, 252) if r_idx % 2 == 0 else RGBColor(255, 255, 255)
                for c_idx, (col_name, val) in enumerate(row.items()):
                    cell = table_shape.cell(r_idx + 1, c_idx)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = row_bg

                    font_color, text = RGBColor(51, 51, 51), "-"
                    if c_idx == 0:
                        text = str(val)
                    elif c_idx == 1:
                        try:
                            if row['Commodity'] == 'Brent':
                                text = f"${float(val):,.2f}"
                            else:
                                text = f"₹{float(val):,.2f}"
                        except:
                            text = str(val)
                    else:
                        if not (pd.isna(val) or str(val).strip() == '-' or val == 0.0):
                            try:
                                val_pct = float(val) * 100
                                text = f"{'+' if val_pct > 0 else ''}{val_pct:.1f}%"
                                font_color = RGBColor(229, 45, 39) if val_pct > 0 else RGBColor(0, 135, 90)
                            except:
                                pass

                    cell.text = text
                    for p in cell.text_frame.paragraphs:
                        p.font.size, p.font.color.rgb = Pt(11), font_color
                        p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
                        if c_idx > 1 and text != "-": p.font.bold = True

    except Exception as e:
        print(f"Warning: Could not build Matrix slide. Error: {e}")

    # 3. Chart Slides
    for page in PAGES_CONFIG:
        chart_path = create_slide_image(page, df_products, df_fx)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_corporate_header(slide, page["title"])

        slide.shapes.add_picture(chart_path, Inches(0.2), Inches(1.3), width=Inches(12.933))

        box_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(5.9), Inches(12.533), Inches(1.1))
        box_bg.fill.solid()
        box_bg.fill.fore_color.rgb = RGBColor(248, 250, 252)
        box_bg.line.color.rgb = RGBColor(226, 232, 240)

        box_accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(5.9), Inches(0.1), Inches(1.1))
        box_accent.fill.solid()
        box_accent.fill.fore_color.rgb = RGBColor(0, 82, 204)
        box_accent.line.fill.background()

        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(6.0), Inches(12.2), Inches(0.9))
        p = txBox.text_frame.add_paragraph()
        p.text = page["insight"]
        p.font.size, p.font.bold, p.font.color.rgb = Pt(13), True, RGBColor(51, 65, 85)

    # 4. Final Appendix Slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_corporate_header(slide, "Data Source & Resource Mapping")

    left, top, width, height = Inches(0.4), Inches(1.6), Inches(12.533), Inches(3.5)
    table_shape = slide.shapes.add_table(len(RESOURCE_LINKS) + 1, 2, left, top, width, height).table
    table_shape.columns[0].width = Inches(3.5)
    table_shape.columns[1].width = Inches(9.033)

    for c_idx, text in enumerate(["Data Source", "Tracked Commodities"]):
        cell = table_shape.cell(0, c_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0, 82, 204)
        for p in cell.text_frame.paragraphs:
            p.font.bold, p.font.color.rgb, p.font.size = True, RGBColor(255, 255, 255), Pt(14)
            p.alignment = PP_ALIGN.LEFT

    for r_idx, (source, components) in enumerate(RESOURCE_LINKS):
        row_bg = RGBColor(248, 250, 252) if r_idx % 2 == 0 else RGBColor(255, 255, 255)
        for c_idx, text in enumerate([source, components]):
            cell = table_shape.cell(r_idx + 1, c_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            for p in cell.text_frame.paragraphs:
                p.font.size, p.font.color.rgb = Pt(13), RGBColor(51, 65, 85)
                if c_idx == 0: p.font.bold = True

    prs.save(PPT_OUTPUT)
    print(f"✓ PowerPoint Presentation saved: {PPT_OUTPUT}")


# ==========================================
# 3. NATIVE PDF GENERATOR ENGINE (WITH TOC)
# ==========================================
def generate_native_pdf():
    print("\nGenerating Executive PDF Report (Cover, TOC + Charts)...")
    try:
        # Safe sheet loading
        try:
            df_products = pd.read_excel(DATA_FILE, sheet_name="Product_Trends")
            df_products['Date'] = pd.to_datetime(df_products['Date'])
        except Exception:
            df_products = pd.DataFrame()

        try:
            df_fx = pd.read_excel(DATA_FILE, sheet_name="Exchange_Rates")
            df_fx['Date'] = pd.to_datetime(df_fx['Date'])
        except Exception:
            df_fx = pd.DataFrame()

        try:
            matrix_df = pd.read_excel(DATA_FILE, sheet_name="Variation_Matrix")
        except Exception:
            matrix_df = pd.DataFrame()

        all_items = []
        if not df_fx.empty:
            for p in sorted(df_fx['Product'].unique()):
                all_items.append({'name': p, 'df': df_fx[df_fx['Product'] == p], 'val_col': 'Price', 'unit': 'INR',
                                  'cat': 'Macro Exchange Rates'})
        if not df_products.empty:
            for p in sorted(df_products['Product'].unique()):
                unit_label = 'USD/BBL' if p == 'Brent' else 'INR/KG'
                all_items.append(
                    {'name': p, 'df': df_products[df_products['Product'] == p], 'val_col': 'Final_INR_KG_Price',
                     'unit': unit_label, 'cat': 'Raw Material Commodities'})

        with PdfPages(PDF_OUTPUT) as pdf:

            # PAGE 1: COVER PAGE
            fig_cover = plt.figure(figsize=(8.27, 11.69))
            fig_cover.patch.set_facecolor('white')
            fig_cover.patches.extend([
                patches.Rectangle((0, 0.85), 1, 0.15, transform=fig_cover.transFigure, clip_on=False,
                                  facecolor='#0052cc'),
                patches.Rectangle((0, 0.0), 1, 0.05, transform=fig_cover.transFigure, clip_on=False,
                                  facecolor='#0f172a')
            ])

            fig_cover.text(0.5, 0.92, "C A M P U S   A C T I V E W E A R", ha='center', va='center', fontsize=26,
                           fontweight='bold', color='white')
            fig_cover.text(0.5, 0.65, "Procurement & Sourcing", ha='center', va='center', fontsize=34,
                           fontweight='bold', color='#0f172a')
            fig_cover.text(0.5, 0.58, "Raw Material Cost & Market Analytics", ha='center', va='center', fontsize=20,
                           color='#334155')
            fig_cover.text(0.5, 0.45, f"Report Generated: {datetime.now().strftime('%d %B %Y')}", ha='center',
                           va='center', fontsize=14, color='#64748b')
            fig_cover.text(0.5, 0.40, "Automated 1-Year Market Intelligence Pipeline", ha='center', va='center',
                           fontsize=12, color='#94a3b8', style='italic')

            pdf.savefig(fig_cover)
            plt.close(fig_cover)

            # PAGE 2: TABLE OF CONTENTS
            fig_toc = plt.figure(figsize=(8.27, 11.69))
            fig_toc.patch.set_facecolor('white')
            fig_toc.text(0.1, 0.9, "Table of Contents", fontsize=24, fontweight='bold', color='#0052cc')
            fig_toc.text(0.1, 0.87, "_______________________________________________________________________________",
                         color='#e2e8f0', fontsize=14)

            start_page, y_pos, col_x, current_category = 4, 0.8, 0.1, ""
            fig_toc.text(col_x, y_pos, "Executive Summary", fontsize=12, fontweight='bold', color='#e52d27')
            y_pos -= 0.04
            fig_toc.text(col_x, y_pos, "Historical Price Variation Matrix", fontsize=9, fontweight='bold',
                         color='#334155')
            fig_toc.text(col_x + 0.50, y_pos, "......... Page 3", ha='right', fontsize=11, color='#64748b')
            y_pos -= 0.03

            for idx, item in enumerate(all_items):
                if idx == 19:
                    col_x, y_pos, current_category = 0.55, 0.8, ""

                if item['cat'] != current_category:
                    y_pos -= 0.02
                    fig_toc.text(col_x, y_pos, item['cat'], fontsize=14, fontweight='bold', color='#e52d27')
                    y_pos -= 0.04
                    current_category = item['cat']

                fig_toc.text(col_x, y_pos, f"{item['name']}", fontsize=11, fontweight='bold', color='#334155')
                fig_toc.text(col_x + 0.35, y_pos, f"......... Page {start_page + idx}", ha='right', fontsize=11,
                             color='#64748b')
                y_pos -= 0.03

            fig_toc.text(0.95, 0.02, "Page 2", ha='right', fontsize=9, color='#64748b')
            pdf.savefig(fig_toc)
            plt.close(fig_toc)

            # PAGE 3: PRICE VARIATION MATRIX TABLE
            if not matrix_df.empty:
                fig_matrix = plt.figure(figsize=(8.27, 11.69))
                fig_matrix.clf()
                fig_matrix.patch.set_facecolor('white')

                fig_matrix.suptitle("Historical Price Variation Matrix", fontsize=18, fontweight='bold',
                                    color='#0f172a', y=0.96)
                ax_matrix = fig_matrix.add_axes([0.02, 0.35, 0.96, 0.55])
                ax_matrix.axis('off')

                display_df = matrix_df.iloc[:, :11].copy()

                def safe_format_price(row):
                    try:
                        x = row['Latest Price']
                        if row['Commodity'] == 'Brent':
                            return f"${float(x):,.2f}"
                        return f"₹{float(x):,.2f}"
                    except Exception:
                        return str(row.get('Latest Price', '-'))

                def safe_format_pct(x):
                    try:
                        if pd.isna(x) or str(x).strip() == '-' or float(x) == 0.0: return "-"
                        return f"{float(x) * 100:+.1f}%"
                    except:
                        return "-"

                display_df['Latest Price'] = display_df.apply(safe_format_price, axis=1)
                for col in display_df.columns[2:]:
                    display_df[col] = display_df[col].apply(safe_format_pct)

                table = ax_matrix.table(cellText=display_df.values, colLabels=display_df.columns, cellLoc='center',
                                        loc='center')
                table.auto_set_font_size(False)
                table.set_fontsize(8)
                table.scale(1, 1.8)

                for (row, col), cell in table.get_celld().items():
                    if row == 0:
                        cell.set_text_props(weight='bold', color='white')
                        cell.set_facecolor('#0052cc')
                    else:
                        cell.set_facecolor('#f8fafc' if row % 2 == 0 else 'white')
                        if col >= 2:
                            val_text = cell.get_text().get_text()
                            if val_text.startswith('+'):
                                cell.set_text_props(color='#e52d27', weight='bold')
                            elif val_text.startswith('-') and len(val_text) > 1:
                                cell.set_text_props(color='#00875a', weight='bold')

                plt.figtext(0.5, 0.30,
                            "Procurement Insight: Red text indicates a price increase (cost risk), while Green text indicates a price decrease (cost savings).",
                            ha='center', fontsize=9, style='italic', color='#475569')
                fig_matrix.text(0.95, 0.02, "Page 3", ha='right', fontsize=9, color='#64748b')
                pdf.savefig(fig_matrix)
                plt.close(fig_matrix)

            # PAGES 4+: INDIVIDUAL COMMODITY DATA PAGES
            for idx, item in enumerate(all_items):
                current_page = start_page + idx
                name, df, val_col, unit = item['name'], item['df'].sort_values('Date', ascending=False), item[
                    'val_col'], item['unit']

                # Graceful Degradation: Create placeholder if data empty
                if df.empty:
                    fig = plt.figure(figsize=(8.27, 11.69))
                    fig.patch.set_facecolor('white')
                    fig.suptitle(f"{name} Market Trend Analysis", fontsize=18, fontweight='bold', color='#0f172a',
                                 y=0.96)
                    fig.text(0.5, 0.5, "⚠️ Data temporarily unavailable for this period", ha='center', va='center',
                             fontsize=16, color='#e52d27', fontweight='bold')
                    fig.text(0.95, 0.02, f"Page {current_page}", ha='right', fontsize=9, color='#64748b')
                    pdf.savefig(fig)
                    plt.close(fig)
                    continue

                fig = plt.figure(figsize=(8.27, 11.69))
                fig.clf()
                fig.patch.set_facecolor('white')
                fig.suptitle(f"{name} Market Trend Analysis", fontsize=18, fontweight='bold', color='#0f172a', y=0.96)

                ax1 = fig.add_axes([0.1, 0.58, 0.8, 0.32])
                df_1yr = df[df['Date'] >= (pd.to_datetime('today') - pd.DateOffset(years=1))]

                monthly = pd.DataFrame()
                if not df_1yr.empty:
                    monthly = df_1yr.set_index('Date').resample('MS').mean(numeric_only=True).reset_index().dropna(
                        subset=[val_col])
                    ax1.plot(monthly['Date'], monthly[val_col], marker='o', linestyle='-', color='#0052cc',
                             linewidth=2.5, markersize=6, label="Monthly Avg")

                    y_max, y_min = monthly[val_col].max(), monthly[val_col].min()
                    y_range = y_max - y_min if y_max != y_min else y_max * 0.1
                    prefix = "$" if unit == "USD/BBL" else ("" if unit == "INR" else "₹")

                    for i_val, (x, y) in enumerate(zip(monthly['Date'], monthly[val_col])):
                        offset_y = y_range * 0.05 if i_val % 2 == 0 else y_range * -0.07
                        format_str = f"{prefix}{y:,.2f}" if unit in ["USD/BBL", "INR"] else f"{prefix}{y:,.0f}"
                        ax1.text(x, y + offset_y, format_str, fontsize=8, ha='center',
                                 va='bottom' if i_val % 2 == 0 else 'top', color='#0f172a', fontweight='bold',
                                 bbox=dict(facecolor='white', edgecolor='none', alpha=0.7, pad=1))

                    ax1.set_title("1-Year Monthly Average Trend", fontsize=12, fontweight='bold', color='#334155',
                                  pad=10)
                    ax1.set_ylabel(f"Price ({unit})", fontsize=10)
                    ax1.xaxis.set_major_locator(mdates.MonthLocator())
                    ax1.xaxis.set_major_formatter(mdates.DateFormatter('%b %y'))
                    plt.setp(ax1.get_xticklabels(), rotation=45, ha='right', fontsize=9)
                    plt.setp(ax1.get_yticklabels(), fontsize=9)
                    ax1.grid(axis='y', linestyle='--', alpha=0.6)
                    ax1.spines['top'].set_visible(False)
                    ax1.spines['right'].set_visible(False)
                    ax1.legend(loc='upper right', frameon=False, fontsize=9)

                ax2 = fig.add_axes([0.1, 0.05, 0.35, 0.45])
                ax2.axis('off')
                ax2.set_title("1-Year Monthly Averages", fontsize=11, fontweight='bold', color='#334155', pad=8)

                if not monthly.empty:
                    monthly_disp = monthly.sort_values('Date', ascending=False).copy()
                    monthly_disp['Month'] = monthly_disp['Date'].dt.strftime('%b %Y')
                    monthly_disp['Price_Str'] = monthly_disp[val_col].apply(
                        lambda x: f"{prefix}{x:,.2f}" if unit in ["USD/BBL", "INR"] else f"{prefix}{x:,.0f}")

                    table_data_m = [['Month', f'Avg Price ({unit})']] + [[row['Month'], row['Price_Str']] for _, row in
                                                                         monthly_disp.iterrows()]
                    table_m = ax2.table(cellText=table_data_m, loc='center', cellLoc='center')
                    table_m.auto_set_font_size(False)
                    table_m.set_fontsize(9)
                    table_m.scale(1, 1.4)

                    for j in range(len(table_data_m[0])):
                        table_m[0, j].set_text_props(weight='bold', color='white', size=10)
                        table_m[0, j].set_facecolor('#0052cc')
                    for r_idx in range(1, len(table_data_m)):
                        for j in range(len(table_data_m[0])):
                            if r_idx % 2 == 0: table_m[r_idx, j].set_facecolor('#f8fafc')

                ax3 = fig.add_axes([0.55, 0.05, 0.35, 0.45])
                ax3.axis('off')

                latest_date = df['Date'].max()
                current_month_df = df[
                    (df['Date'].dt.year == latest_date.year) & (df['Date'].dt.month == latest_date.month)].sort_values(
                    'Date', ascending=False)
                ax3.set_title(f"Daily Data ({latest_date.strftime('%b %Y')})", fontsize=11, fontweight='bold',
                              color='#334155', pad=8)

                display_df = current_month_df.head(28).copy()
                if not display_df.empty:
                    display_df['Date_Str'] = display_df['Date'].dt.strftime('%d-%b-%y')
                    display_df['Price_Str'] = display_df[val_col].apply(
                        lambda x: f"{prefix}{x:,.2f}" if unit in ["USD/BBL", "INR"] else f"{prefix}{x:,.0f}")

                    table_data_d = [['Date', f'Daily Price ({unit})']] + [[row['Date_Str'], row['Price_Str']] for _, row
                                                                          in display_df.iterrows()]
                    table_d = ax3.table(cellText=table_data_d, loc='center', cellLoc='center')
                    table_d.auto_set_font_size(False)
                    table_d.set_fontsize(9 if len(table_data_d) < 18 else 8)
                    table_d.scale(1, 1.4 if len(table_data_d) < 18 else 1.1)

                    for j in range(len(table_data_d[0])):
                        table_d[0, j].set_text_props(weight='bold', color='white', size=9)
                        table_d[0, j].set_facecolor('#0052cc')
                    for r_idx in range(1, len(table_data_d)):
                        for j in range(len(table_data_d[0])):
                            if r_idx % 2 == 0: table_d[r_idx, j].set_facecolor('#f8fafc')

                fig.text(0.95, 0.02, f"Page {current_page}", ha='right', fontsize=9, color='#64748b')
                pdf.savefig(fig)
                plt.close(fig)

        print(f"✓ Native PDF Report successfully generated: {PDF_OUTPUT}")
    except Exception as e:
        print(f"\nCould not generate native PDF. Error: {e}")


if __name__ == "__main__":
    pre_flight_checks()
    build_presentation()
    generate_native_pdf()
    post_flight_validation()